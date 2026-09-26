import io
import os
import re
import unicodedata
from typing import Optional, Tuple
from fastapi import HTTPException
import pypdf
import docx


SUPPORTED_EXTENSIONS = {".txt", ".md", ".pdf", ".docx"}


def validate_extracted_content(content: Optional[str], document_name: Optional[str] = None) -> str:
    """
    Pre-generation Extraction Validation Layer.
    Deterministically validates that source/extracted content is present, meaningful,
    well-encoded, and uncorrupted before prompt construction and Gemini API invocation.

    Raises HTTPException(status_code=400, detail=...) if validation fails.
    Returns the cleaned, validated text string if successful.
    """
    doc_ref = f" from '{document_name}'" if document_name else ""

    # Rule 1: Presence & non-empty
    if content is None:
        raise HTTPException(
            status_code=400,
            detail=f"Extraction validation failed: No content could be extracted{doc_ref}."
        )

    text = content.strip()
    if not text:
        raise HTTPException(
            status_code=400,
            detail=f"Extraction validation failed: Extracted content{doc_ref} is empty or whitespace-only."
        )

    # Allow visual/video source placeholders (e.g. "[Visual Source: chart.png]", "[Video Source: demo.mp4]")
    if text.startswith("[Visual Source:") or text.startswith("[Video Source:") or text.startswith("[Multimodal Analysis"):
        return text

    # Rule 2: Minimum viable length
    if len(text) < 5 or len(text.split()) < 2:
        alnum_chars = [c for c in text if c.isalnum()]
        if len(alnum_chars) < 4:
            raise HTTPException(
                status_code=400,
                detail=f"Extraction validation failed: Extracted text{doc_ref} is too short to be viable source content ({len(text)} characters)."
            )

    # Rule 3: Encoding & Null bytes & Control characters
    # Null bytes indicate raw binary data or catastrophic decoding
    if "\x00" in text:
        raise HTTPException(
            status_code=400,
            detail=f"Extraction validation failed: Extracted text{doc_ref} contains binary null bytes (corrupted extraction)."
        )

    # Unicode replacement character '\ufffd' (indicates failed byte decoding)
    replacement_count = text.count("\ufffd")
    if replacement_count > 3 or (len(text) > 0 and (replacement_count / len(text)) > 0.02):
        raise HTTPException(
            status_code=400,
            detail=f"Extraction validation failed: Extracted text{doc_ref} contains corrupted unicode decoding artifacts."
        )

    # Dangerous control characters (excluding newline \n, carriage return \r, tab \t)
    control_chars = [c for c in text if ord(c) < 32 and c not in ("\t", "\n", "\r")]
    if len(control_chars) > 0:
        if len(control_chars) > 3 or (len(control_chars) / len(text)) > 0.02:
            raise HTTPException(
                status_code=400,
                detail=f"Extraction validation failed: Extracted text{doc_ref} contains unreadable control characters."
            )

    # Rule 4: Obvious extraction failure / error markers
    text_lower = text.lower()
    error_markers = [
        "error: could not extract text",
        "no readable text could be extracted",
        "failed to extract text",
        "failed to read document",
        "unsupported file format",
        "pdfreaderror",
        "badzipfile",
        "traceback (most recent call last)"
    ]
    for marker in error_markers:
        if marker in text_lower and len(text) < 300:
            raise HTTPException(
                status_code=400,
                detail=f"Extraction validation failed: Extracted text{doc_ref} contains an extraction error message instead of source content."
            )

    # Rule 5: Suspiciously poor / garbage extraction
    # Fully Unicode-aware: Letters (L), Numbers (N), Combining Marks (M, e.g. Hindi/Arabic/accent matras),
    # Punctuation (P), Symbols (S), and Whitespace (Z/isspace) are valid human characters.
    # Control/Unassigned/Surrogate characters (category C) indicate binary corruption.
    readable_count = sum(
        1 for c in text
        if c.isspace() or unicodedata.category(c)[0] in ("L", "N", "M", "P", "S")
    )
    if len(text) > 20 and (readable_count / len(text)) < 0.70:
        raise HTTPException(
            status_code=400,
            detail=f"Extraction validation failed: Extracted text{doc_ref} appears to be corrupted or unreadable binary garbage."
        )

    # Check repeating single character spam (e.g. "xxxxxxxxxxxxxxxxxxxx" or "....................")
    if len(text) >= 30 and len(set(text.replace(" ", "").replace("\n", ""))) <= 2:
        raise HTTPException(
            status_code=400,
            detail=f"Extraction validation failed: Extracted text{doc_ref} contains meaningless repetitive character sequences."
        )

    return text


def extract_text_from_document(filename: str, file_bytes: bytes) -> str:
    """
    Extracts readable text from supported document formats (.txt, .md, .pdf, .docx).
    Raises HTTPException with clear error details if the format is unsupported,
    corrupted, or contains no extractable text.
    """
    if not filename:
        raise HTTPException(
            status_code=400,
            detail="Invalid request: Missing document filename."
        )

    _, ext = os.path.splitext(filename.lower())

    if ext not in SUPPORTED_EXTENSIONS:
        raise HTTPException(
            status_code=400,
            detail=(
                f"Unsupported document format '{ext}'. "
                f"Supported formats are: {', '.join(sorted(SUPPORTED_EXTENSIONS))}."
            )
        )

    if not file_bytes or len(file_bytes) == 0:
        raise HTTPException(
            status_code=400,
            detail="The uploaded file is empty."
        )

    extracted_text = ""

    try:
        if ext in {".txt", ".md"}:
            try:
                extracted_text = file_bytes.decode("utf-8")
            except UnicodeDecodeError:
                extracted_text = file_bytes.decode("latin-1")

        elif ext == ".pdf":
            pdf_stream = io.BytesIO(file_bytes)
            reader = pypdf.PdfReader(pdf_stream)

            if reader.is_encrypted:
                try:
                    # Attempt decrypt with empty password if standard unencrypted container
                    reader.decrypt("")
                except Exception:
                    raise HTTPException(
                        status_code=400,
                        detail="The uploaded PDF is password-protected and cannot be read."
                    )

            page_texts = []
            for i, page in enumerate(reader.pages):
                page_text = page.extract_text() or ""
                if page_text.strip():
                    page_texts.append(page_text.strip())

            extracted_text = "\n\n".join(page_texts)

        elif ext == ".docx":
            docx_stream = io.BytesIO(file_bytes)
            doc = docx.Document(docx_stream)

            paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

            # Also check table content
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        paragraphs.append(row_text)

            extracted_text = "\n\n".join(paragraphs)

    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(
            status_code=422,
            detail=f"Failed to extract text from {ext} file: {str(e)}"
        )

    cleaned_text = extracted_text.strip()
    if not cleaned_text:
        raise HTTPException(
            status_code=400,
            detail=(
                f"No readable text could be extracted from '{filename}'. "
                "If this is a scanned document or image, please provide a document with selectable text."
            )
        )

    # Validate extracted text through the Extraction Validation Layer
    return validate_extracted_content(cleaned_text, document_name=filename)
